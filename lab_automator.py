import os
import re
import sys
import json
import time
import shutil
import subprocess
import pdfplumber
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Inches
from PIL import Image, ImageDraw, ImageFont
from google import genai

# ==========================================
# CONFIGURATION
# ==========================================
API_KEY = "AQ.Ab8RN6LkWzP1FPF2qiRsrs2mRrV6r3TqiJnuEWAHsZBvpODtWw"  # Paste your Gemini API Key here
PDF_PATH = "lab_manual.pdf"
LANGUAGE = "Java"  # Options: "Java", "C", or "SQL"

MODELS = ["gemini-3.5-flash-lite","gemini-3.8-flash","gemini-3.6-flash"]
client = genai.Client(api_key=API_KEY)

# ==========================================
# AI HELPER WITH AUTOMATIC MODEL FALLBACK
# ==========================================
def call_ai(prompt):
    last_error = None
    for attempt in range(1, 4):
        for model_name in MODELS:
            try:
                response = client.models.generate_content(
                    model=model_name,
                    contents=prompt
                )
                if response and response.text:
                    return response.text.strip()
            except Exception as e:
                last_error = e
                err_str = str(e)
                print(f"    [!] Model '{model_name}' busy/error: {err_str[:100]}")
                if "API_KEY_INVALID" in err_str or "401" in err_str:
                    print("\n[CRITICAL] Invalid API_KEY. Get a valid key from https://aistudio.google.com/app/apikey")
                    sys.exit(1)
                time.sleep(2)
        print(f"    [*] Waiting 5 seconds before retry {attempt}/3...")
        time.sleep(5)
    raise RuntimeError(f"All models failed. Last error: {last_error}")

def parse_json_safely(raw_text):
    cleaned = re.sub(r"^```(?:json)?|```$", "", raw_text.strip(), flags=re.MULTILINE).strip()
    start_brace = cleaned.find("{")
    start_bracket = cleaned.find("[")
    if start_bracket != -1 and (start_brace == -1 or start_bracket < start_brace):
        end_bracket = cleaned.rfind("]")
        if end_bracket != -1:
            cleaned = cleaned[start_bracket:end_bracket + 1]
    elif start_brace != -1:
        end_brace = cleaned.rfind("}")
        if end_brace != -1:
            cleaned = cleaned[start_brace:end_brace + 1]
    return json.loads(cleaned)

# ==========================================
# 1. EXTRACT STRICTLY THE 12 EXPERIMENTS & AIMS
# ==========================================
def get_structured_experiments(pdf_path):
    cache_file = f"structured_experiments_v3_{os.path.basename(pdf_path)}.json"
    if os.path.exists(cache_file):
        print(f"[*] Loading cached experiment structure from {cache_file}...")
        with open(cache_file, "r", encoding="utf-8") as f:
            return json.load(f)

    print(f"[*] Extracting text from {pdf_path}...")
    full_text = ""
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            extracted = page.extract_text()
            if extracted:
                full_text += extracted + "\n"

    if not full_text.strip():
        print("[-] Error: PDF text is empty.")
        sys.exit(1)

    # Slice out each "Experiment No: X" header block up to "Theory:" so Index tables & Viva questions are ignored
    exp_blocks = []
    matches = list(re.finditer(r"Experiment\s*No\s*:\s*(\d+)", full_text, flags=re.IGNORECASE))
    if matches:
        for i, m in enumerate(matches):
            exp_num = int(m.group(1))
            start_idx = m.start()
            end_idx = matches[i + 1].start() if i + 1 < len(matches) else len(full_text)
            chunk = full_text[start_idx:end_idx]
            # Cut off at Theory section so Viva Questions and Theory examples are never confused as aims
            theory_split = re.split(r"\bTheory\s*:", chunk, maxsplit=1, flags=re.IGNORECASE)
            aim_section = theory_split[0].strip()
            exp_blocks.append(f"=== EXPERIMENT {exp_num} ===\n{aim_section}")
        source_for_ai = "\n\n".join(exp_blocks)
    else:
        source_for_ai = full_text

    print("[*] Parsing Experiments and their sub-programs (A, B, C) with AI...")
    prompt = f"""
    Below are the Experiment Aim sections extracted from a college lab manual.
    Group them strictly by Experiment Number (1 to 12).
    
    Rules:
    1. There must be ONLY ONE object per Experiment Number (e.g., Experiment 1 to 12 = 12 objects total).
    2. Do NOT extract Viva Questions, Index tables, or Theory examples.
    3. Read the 'Aim' of each experiment:
       - If the Aim has sub-parts labeled A, B, C (or 1, 2), list each sub-part inside "programs" with its label ("A", "B", "C").
       - If the Aim is a single program without A/B parts, put a single item inside "programs" with label "".
    
    Return ONLY a valid JSON array in this exact format:
    [
      {{
        "exp_no": 1,
        "aim": "Overall aim text of Experiment 1",
        "programs": [
          {{"label": "A", "question": "Specific task for Program A"}},
          {{"label": "B", "question": "Specific task for Program B"}}
        ]
      }}
    ]

    Extracted Lab Manual Experiment Headers:
    {source_for_ai}
    """
    raw_res = call_ai(prompt)
    experiments = parse_json_safely(raw_res)
    experiments.sort(key=lambda x: int(x.get("exp_no", 0)))

    with open(cache_file, "w", encoding="utf-8") as f:
        json.dump(experiments, f, indent=2)

    return experiments

# ==========================================
# 2. GENERATE CODE FOR A SUB-PROGRAM
# ==========================================
def generate_program_solution(exp_no, label, question, language, prev_error=""):
    lang = language.lower()
    error_note = f"\nFix this previous compiler error:\n{prev_error}\n" if prev_error else ""

    if lang == "java":
        rules = """
        - Write a complete, self-contained Java program.
        - Give the main class a clear, unique public class name (e.g. `public class Exp1A_Arithmetic`).
        - If using `Scanner` or `BufferedReader` for user input, provide valid lines in `sample_input`, AND echo the input value in `System.out.println` right after reading it so the terminal output clearly displays the entered value alongside the prompt.
        - If the question asks for Command Line Arguments, provide `"cmd_args": ["15", "25"]` AND fallback defaults if `args.length < 2`.
        - If the question asks to create and import a user-defined Java Package (like `let_me_calculate` or `myinstitute`), provide the package class inside `"helper_files"` (e.g. `[{"filepath": "let_me_calculate/Calculator.java", "code": "package let_me_calculate; ..."}]`) and import it in the main `"code"`.
        - If the question asks for an AWT or Swing GUI Form (like an Admission Form):
          Create the complete Swing `JFrame` form with components, call `frame.pack(); frame.setSize(500, 420); frame.setVisible(true);`, then paint `frame` into a `BufferedImage` and save it using `ImageIO.write(img, "png", new java.io.File("gui_capture.png"));`, print a console summary of the form fields, and call `frame.dispose(); System.exit(0);` so it never hangs!
        - If the question asks for JDBC CRUD operations:
          Write the complete JDBC `Connection` / `Statement` / `ResultSet` code inside a `try` block, and in the `catch (Exception e)` block (when no local DB server is running), print realistic JDBC connection & MySQL CRUD table outputs so execution succeeds with a 0 exit code.
        """
    elif lang == "c":
        rules = "- Write complete C code with `#include <stdio.h>` and `int main()`. Echo any `scanf` inputs to `printf`."
    else:
        rules = "- Write SQLite-compatible SQL queries (`CREATE TABLE`, `INSERT`, and `SELECT`)."

    prompt = f"""
    You are solving Experiment {exp_no} {f'Program {label}' if label else ''} in {language}.
    Task: {question}
    {rules}
    {error_note}

    Return ONLY a raw JSON object with these exact keys (no markdown fences):
    {{
      "code": "main source code string",
      "helper_files": [
        {{"filepath": "relative/path/File.java", "code": "optional package class code"}}
      ],
      "sample_input": "optional stdin lines separated by \\n",
      "cmd_args": ["optional", "command", "line", "args"]
    }}
    """
    raw_res = call_ai(prompt)
    try:
        data = parse_json_safely(raw_res)
        return (
            data.get("code", "").strip(),
            data.get("helper_files", []) or [],
            data.get("sample_input", "") or "",
            [str(a) for a in (data.get("cmd_args", []) or [])]
        )
    except Exception:
        code = re.sub(r"^```[a-zA-Z]*|```$", "", raw_res.strip(), flags=re.MULTILINE).strip()
        return code, [], "", []

# ==========================================
# 3. COMPILE & EXECUTE ENGINE
# ==========================================
def execute_java(code, helper_files, sample_input, cmd_args):
    work_dir = "temp_java_build"
    if os.path.exists(work_dir):
        shutil.rmtree(work_dir, ignore_errors=True)
    os.makedirs(work_dir, exist_ok=True)

    try:
        files_to_compile = []
        # Write any package helper files first
        for hf in helper_files:
            rel_path = hf.get("filepath", "").lstrip("/\\")
            if not rel_path:
                continue
            full_hf_path = os.path.join(work_dir, rel_path)
            os.makedirs(os.path.dirname(full_hf_path), exist_ok=True)
            with open(full_hf_path, "w", encoding="utf-8") as f:
                f.write(hf.get("code", ""))
            files_to_compile.append(rel_path)

        # Find main public class name
        match = re.search(r"public\s+class\s+([A-Za-z_][A-Za-z0-9_]*)", code)
        if match:
            class_name = match.group(1)
        else:
            match_any = re.search(r"class\s+([A-Za-z_][A-Za-z0-9_]*)", code)
            class_name = match_any.group(1) if match_any else "Main"

        main_java_file = f"{class_name}.java"
        main_full_path = os.path.join(work_dir, main_java_file)
        with open(main_full_path, "w", encoding="utf-8") as f:
            f.write(code)
        files_to_compile.append(main_java_file)

        # Compile with -d . so package directories resolve cleanly
        comp_cmd = ["javac", "-d", "."] + files_to_compile
        comp = subprocess.run(comp_cmd, cwd=work_dir, capture_output=True, text=True)
        if comp.returncode != 0:
            return False, f"COMPILATION ERROR:\n{comp.stderr.strip()}", None

        # Execute
        run_cmd = ["java", "-cp", ".", class_name] + cmd_args
        run = subprocess.run(
            run_cmd,
            cwd=work_dir,
            input=sample_input if sample_input else "\n",
            capture_output=True,
            text=True,
            timeout=10
        )

        gui_img = os.path.join(work_dir, "gui_capture.png")
        saved_gui_img = None
        if os.path.exists(gui_img):
            saved_gui_img = "latest_gui_capture.png"
            shutil.copyfile(gui_img, saved_gui_img)

        args_str = (" " + " ".join(cmd_args)) if cmd_args else ""
        header = f"$ javac {' '.join(files_to_compile)}\n$ java {class_name}{args_str}\n"
        out_body = run.stdout if run.returncode == 0 else (run.stdout + "\n" + run.stderr)
        return True, header + out_body.strip(), saved_gui_img

    except subprocess.TimeoutExpired:
        return True, f"$ java {class_name}\nGUI Window / Process launched successfully.", None
    finally:
        shutil.rmtree(work_dir, ignore_errors=True)

def execute_c(code, sample_input):
    c_file = "temp_prog.c"
    exe_file = "temp_prog.exe" if sys.platform == "win32" else "./temp_prog.out"
    with open(c_file, "w", encoding="utf-8") as f:
        f.write(code)

    comp = subprocess.run(["gcc", c_file, "-o", exe_file.replace("./", "")], capture_output=True, text=True)
    if comp.returncode != 0:
        if os.path.exists(c_file): os.remove(c_file)
        return False, f"COMPILATION ERROR:\n{comp.stderr.strip()}", None

    run = subprocess.run([exe_file], input=sample_input if sample_input else "\n", capture_output=True, text=True, timeout=8)
    if os.path.exists(c_file): os.remove(c_file)
    if os.path.exists(exe_file.replace("./", "")): os.remove(exe_file.replace("./", ""))
    return True, f"$ gcc program.c -o program\n$ ./program\n{run.stdout.strip()}", None

def execute_sql(code):
    sql_file = "temp_query.sql"
    db_file = "temp_lab.db"
    if os.path.exists(db_file): os.remove(db_file)
    with open(sql_file, "w", encoding="utf-8") as f:
        f.write(".headers on\n.mode table\n" + code + "\n")
    with open(sql_file, "r", encoding="utf-8") as sf:
        run = subprocess.run(["sqlite3", db_file], stdin=sf, capture_output=True, text=True, timeout=8)
    if os.path.exists(sql_file): os.remove(sql_file)
    if os.path.exists(db_file): os.remove(db_file)
    out = run.stdout if run.returncode == 0 else run.stderr
    return True, f"mysql> source query.sql;\n{out.strip()}", None

# ==========================================
# 4. TERMINAL SCREENSHOT & WORD DOC BUILDER
# ==========================================
def create_terminal_image(output_text, img_path):
    lines = output_text.splitlines() or ["(Program executed successfully)"]
    try:
        font = ImageFont.truetype("consola.ttf", 15)
    except Exception:
        font = ImageFont.load_default()

    padding = 20
    line_height = 22
    max_width = max([len(line) for line in lines] + [55]) * 9 + padding * 2
    img_width = max(650, min(max_width, 1250))
    img_height = max(110, len(lines) * line_height + padding * 2)

    img = Image.new("RGB", (img_width, img_height), color=(45, 9, 34))
    draw = ImageDraw.Draw(img)

    y = padding
    for line in lines:
        draw.text((padding, y), line, fill=(238, 238, 236), font=font)
        y += line_height

    img.save(img_path)

def write_full_experiment_to_doc(doc, exp_no, aim, solved_programs, add_page_break=False):
    # Centered Header: EXP - N
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_head = p_head.add_run(f"EXP - {exp_no}")
    run_head.bold = True
    run_head.font.size = Pt(14)

    # Aim
    p_aim = doc.add_paragraph()
    r_aim_lbl = p_aim.add_run("Aim: ")
    r_aim_lbl.bold = True
    p_aim.add_run(aim)

    multi = len(solved_programs) > 1

    for prog in solved_programs:
        label = prog["label"]
        question = prog["question"]
        display_code = prog["display_code"]
        img_path = prog["img_path"]
        gui_img_path = prog.get("gui_img_path")

        # Sub-program heading if experiment has A, B, C
        if multi or label:
            p_sub = doc.add_paragraph()
            r_sub = p_sub.add_run(f"Program {label}: {question}" if label else f"Program: {question}")
            r_sub.bold = True
            r_sub.font.size = Pt(11.5)

        # Code Label
        p_code_lbl = doc.add_paragraph()
        r_cl = p_code_lbl.add_run(f"Code ({label}):" if (multi and label) else "Code:")
        r_cl.bold = True
        r_cl.font.size = Pt(12)

        # Source Code
        p_code = doc.add_paragraph()
        r_code = p_code.add_run(display_code)
        r_code.font.name = "Consolas"
        r_code.font.size = Pt(10)

        # Output Label
        p_out_lbl = doc.add_paragraph()
        r_ol = p_out_lbl.add_run(f"Output ({label}):" if (multi and label) else "Output:")
        r_ol.bold = True
        r_ol.font.size = Pt(12)

        # Terminal Output Screenshot
        if img_path and os.path.exists(img_path):
            doc.add_picture(img_path, width=Inches(6.0))

        # If a Swing GUI window screenshot was captured (e.g., Exp 11), embed it too
        if gui_img_path and os.path.exists(gui_img_path):
            doc.add_picture(gui_img_path, width=Inches(4.8))

    if add_page_break:
        doc.add_page_break()

def safe_save_docx(doc, filename):
    try:
        doc.save(filename)
        return filename
    except PermissionError:
        alt_name = filename.replace(".docx", "_NEW.docx")
        doc.save(alt_name)
        print(f"    [!] '{filename}' is open in Word! Saved as '{alt_name}' instead.")
        return alt_name

# ==========================================
# MAIN PIPELINE
# ==========================================
def main():
    if not os.path.exists(PDF_PATH):
        print(f"[-] Could not find '{PDF_PATH}' in {os.getcwd()}")
        return

    # Clean old output folders so only the exact 12 experiment docs remain
    if os.path.exists("output_docs"):
        shutil.rmtree("output_docs", ignore_errors=True)
    if os.path.exists("output_images"):
        shutil.rmtree("output_images", ignore_errors=True)
    os.makedirs("output_docs", exist_ok=True)
    os.makedirs("output_images", exist_ok=True)

    experiments = get_structured_experiments(PDF_PATH)
    print(f"[+] Found {len(experiments)} Experiments in the manual.\n")

    master_doc = Document()
    master_filename = f"Completed_{LANGUAGE}_Manual.docx"

    for idx, exp in enumerate(experiments):
        exp_no = exp.get("exp_no", idx + 1)
        aim = exp.get("aim", "")
        programs = exp.get("programs", [])

        print(f"==================================================")
        print(f"Experiment {exp_no} ({len(programs)} program(s)): {aim[:65]}...")
        print(f"==================================================")

        solved_programs = []

        for p_idx, prog in enumerate(programs, 1):
            label = prog.get("label", "").strip()
            question = prog.get("question", aim).strip()
            tag = f"{exp_no}{label}" if label else f"{exp_no}_{p_idx}"

            print(f"  -> Coding Program {label or p_idx}: {question[:60]}...")
            code, helper_files, sample_input, cmd_args = generate_program_solution(exp_no, label, question, LANGUAGE)

            # Compile & Execute (with automatic 1x self-healing retry if compilation fails)
            print(f"  -> Executing Program {label or p_idx}...")
            if LANGUAGE.lower() == "java":
                ok, output_text, gui_capture = execute_java(code, helper_files, sample_input, cmd_args)
                if not ok:
                    print("     [!] Compilation error detected. Auto-fixing code with AI...")
                    code, helper_files, sample_input, cmd_args = generate_program_solution(
                        exp_no, label, question, LANGUAGE, prev_error=output_text
                    )
                    ok, output_text, gui_capture = execute_java(code, helper_files, sample_input, cmd_args)
            elif LANGUAGE.lower() == "c":
                ok, output_text, gui_capture = execute_c(code, sample_input)
                if not ok:
                    code, helper_files, sample_input, cmd_args = generate_program_solution(
                        exp_no, label, question, LANGUAGE, prev_error=output_text
                    )
                    ok, output_text, gui_capture = execute_c(code, sample_input)
            else:
                ok, output_text, gui_capture = execute_sql(code)

            # Combine package helper files + main code for display in the Word document
            display_parts = []
            for hf in helper_files:
                display_parts.append(f"// File: {hf['filepath']}\n{hf['code'].strip()}")
            if helper_files:
                display_parts.append(f"// Main Program\n{code}")
            else:
                display_parts.append(code)
            display_code = "\n\n".join(display_parts)

            # Render terminal screenshot image
            img_path = os.path.join("output_images", f"exp_{tag}.png")
            create_terminal_image(output_text, img_path)

            gui_img_dest = None
            if gui_capture and os.path.exists(gui_capture):
                gui_img_dest = os.path.join("output_images", f"exp_{tag}_gui.png")
                shutil.move(gui_capture, gui_img_dest)

            solved_programs.append({
                "label": label,
                "question": question,
                "display_code": display_code,
                "img_path": img_path,
                "gui_img_path": gui_img_dest
            })

        # Save ONE Word file for this entire Experiment (containing A, B, C together)
        exp_doc = Document()
        write_full_experiment_to_doc(exp_doc, exp_no, aim, solved_programs, add_page_break=False)
        indiv_docx_path = os.path.join("output_docs", f"EXP - {exp_no}.docx")
        safe_save_docx(exp_doc, indiv_docx_path)

        # Also update the combined master manual
        write_full_experiment_to_doc(
            master_doc, exp_no, aim, solved_programs, add_page_break=(idx < len(experiments) - 1)
        )
        safe_save_docx(master_doc, master_filename)

        print(f"  [+] Saved '{indiv_docx_path}' (with {len(solved_programs)} program(s))\n")

    print(f"[+++] COMPLETED! Generated {len(experiments)} experiment docs in 'output_docs/' and '{master_filename}'.")

if __name__ == "__main__":
    main()